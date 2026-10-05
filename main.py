import io
import json
import re

import pymupdf
from docx import Document
from fastapi import FastAPI, UploadFile, File, Form
from fastapi.responses import HTMLResponse
from sentence_transformers import SentenceTransformer, util

app = FastAPI()
model = SentenceTransformer("all-MiniLM-L6-v2")

with open("skills.json", "r", encoding="utf-8") as f:
    skills = json.load(f)

recommendation_messages = {
    "python": "Add Python projects or practical experience to demonstrate your Python skills.",
    "java": "Add a Java project or mention Java-based coursework or experience.",
    "c++": "Add a C++ project or problem-solving experience using C++.",
    "javascript": "Add JavaScript projects that demonstrate practical frontend or backend development.",
    "typescript": "Consider adding a TypeScript project to demonstrate your experience with typed JavaScript.",
    "html": "Make sure your projects demonstrate practical HTML development.",
    "css": "Add a project showing practical CSS and responsive web design skills.",
    "react": "Consider adding a React project and mention the features you implemented.",
    "node.js": "Add a Node.js backend project to demonstrate server-side development.",
    "sql": "Add a project involving databases and SQL queries.",
    "git": "Mention your Git/GitHub workflow and contributions to projects.",
    "docker": "Consider adding Docker to a project and mention how you used containers.",
    "mongodb": "Add a project that uses MongoDB and describe how you stored and retrieved data.",
    "postgresql": "Consider adding a project using PostgreSQL and mention the database operations you performed.",
    "aws": "Consider deploying a project using AWS and mentioning the services you used.",
    "machine learning": "Add a machine learning project and briefly describe the model, dataset, and results.",
    "fastapi": "Add a FastAPI project and mention the APIs or backend features you built.",
    "django": "Consider adding a Django project demonstrating backend development.",
}


def skill_found(skill, text):
    if skill in ["c++", "node.js"]:
        return skill in text
    return re.search(r"\b" + re.escape(skill) + r"\b", text) is not None


@app.get("/", response_class=HTMLResponse)
def home():
    with open("templates/index.html", "r", encoding="utf-8") as file:
        html_content = file.read()
    return html_content


@app.post("/upload")
def upload_resume(
    file: UploadFile = File(...),
    job_description: str = Form(...)
):
    if not job_description.strip():
        return {"error": "Please enter a job description."}

    filename = (file.filename or "").lower()
    file_bytes = file.file.read()
    text = ""

    if filename.endswith(".pdf"):
        try:
            pdf = pymupdf.open(stream=file_bytes, filetype="pdf")
        except Exception:
            return {"error": "Unable to read this PDF. Please upload a valid PDF resume."}
        for page in pdf:
            text += page.get_text()

    elif filename.endswith(".docx"):
        try:
            doc = Document(io.BytesIO(file_bytes))
        except Exception:
            return {"error": "Unable to read this Word file. Please upload a valid DOCX resume."}
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + "\n"

    else:
        return {"error": "Please upload a PDF or DOCX file."}

    if not text.strip():
        return {"error": "No readable text found in this file. Please upload a text-based resume."}

    resume_lower = text.lower()
    jd_lower = job_description.lower()

    # --- AI semantic score ---
    resume_embedding = model.encode(text, convert_to_tensor=True)
    jd_embedding = model.encode(job_description, convert_to_tensor=True)
    semantic_score = util.cos_sim(resume_embedding, jd_embedding).item()
    semantic_score = max(round(semantic_score * 100, 2), 0)

    # --- Resume checks ---
    email_found = re.search(r"[\w\.-]+@[\w\.-]+\.\w+", text) is not None
    phone_found = re.search(r"(?<!\d)(?:\+91[\s-]?)?[6-9]\d{4}[\s-]?\d{5}(?!\d)", text) is not None
    education_found = (
        "education" in resume_lower
        or "academic" in resume_lower
        or "qualification" in resume_lower
    )
    experience_found = (
        "experience" in resume_lower
        or "employment" in resume_lower
        or "work history" in resume_lower
        or "internship" in resume_lower
    )
    projects_found = "project" in resume_lower
    skills_section_found = "skills" in resume_lower

    resume_checks = [
        email_found,
        phone_found,
        education_found,
        experience_found,
        projects_found,
        skills_section_found,
    ]
    resume_completeness = round((sum(resume_checks) / len(resume_checks)) * 100, 2)

    # --- Skill matching ---
    matched_skills = []
    missing_skills = []
    for skill in skills:
        if skill_found(skill, jd_lower):
            if skill_found(skill, resume_lower):
                matched_skills.append(skill)
            else:
                missing_skills.append(skill)

    total_required = len(matched_skills) + len(missing_skills)

    if total_required > 0:
        match_percentage = round((len(matched_skills) / total_required) * 100, 2)
    else:
        match_percentage = 0

    # --- Overall score ---
    if total_required > 0:
        overall_score = round(
            match_percentage * 0.5 + semantic_score * 0.3 + resume_completeness * 0.2, 2
        )
    else:
        overall_score = round(semantic_score * 0.6 + resume_completeness * 0.4, 2)

    if overall_score >= 70:
        overall_label = "Strong"
    elif overall_score >= 50:
        overall_label = "Average"
    else:
        overall_label = "Weak"

    # --- Recommendations ---
    recommendations = []

    for skill in missing_skills:
        recommendations.append(
            recommendation_messages.get(
                skill,
                "Consider adding " + skill + " experience or projects to your resume."
            )
        )

    if not email_found:
        recommendations.append("Add a professional email address to your resume.")

    if not phone_found:
        recommendations.append("Add a phone number to your resume.")

    if not education_found:
        recommendations.append(
            "Add an Education section with your degree, university, and graduation year."
        )

    if not experience_found:
        recommendations.append(
            "Add an Experience section if you have internships, work experience, or relevant practical experience."
        )

    if not projects_found:
        recommendations.append(
            "Add a Projects section with relevant projects, technologies used, and your contributions."
        )

    if not skills_section_found:
        recommendations.append(
            "Add a dedicated Skills section listing relevant programming languages, frameworks, databases, and tools."
        )

    return {
        "matched_skills": matched_skills,
        "missing_skills": missing_skills,
        "match_percentage": match_percentage,
        "semantic_score": semantic_score,
        "overall_score": overall_score,
        "overall_label": overall_label,
        "recommendations": recommendations,
        "email_found": email_found,
        "phone_found": phone_found,
        "education_found": education_found,
        "experience_found": experience_found,
        "projects_found": projects_found,
        "skills_section_found": skills_section_found,
        "resume_completeness": resume_completeness,
    }